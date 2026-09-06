"""Excel export utilities using openpyxl."""

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from flask import send_file


def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column].width = min(max_length + 2, 40)


def _send_workbook(wb, filename):
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


def export_products(products):
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    headers = [
        "Product", "Category", "Brand", "Unit", "Purchase Price",
        "Selling Price", "Stock", "Minimum Stock", "Batch Number", "Expiry Date",
    ]
    ws.append(headers)
    for p in products:
        ws.append([
            p.name,
            p.category.name if p.category else "",
            p.brand.name if p.brand else "",
            p.unit.name if p.unit else "",
            p.purchase_price,
            p.selling_price,
            p.stock,
            p.minimum_stock,
            p.batch_number or "",
            p.expiry_date.strftime("%Y-%m-%d") if p.expiry_date else "",
        ])
    _style_header(ws)
    return _send_workbook(wb, "products.xlsx")


def export_customers(customers):
    from models.bill import Bill
    from models.payment import Payment
    from sqlalchemy import func
    from extensions import db

    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    headers = [
        "Customer", "Phone", "Address", "Opening Balance",
        "Sales", "Payments", "Current Balance",
    ]
    ws.append(headers)
    for c in customers:
        sales = db.session.query(
            func.coalesce(func.sum(Bill.grand_total), 0)
        ).filter(Bill.customer_id == c.id).scalar()
        payments = db.session.query(
            func.coalesce(func.sum(Payment.amount), 0)
        ).filter(Payment.customer_id == c.id, Payment.payment_type == "received").scalar()
        ws.append([
            c.name, c.phone or "", c.address or "",
            c.opening_balance, sales, payments, c.balance,
        ])
    _style_header(ws)
    return _send_workbook(wb, "customers.xlsx")


def export_suppliers(suppliers):
    wb = Workbook()
    ws = wb.active
    ws.title = "Suppliers"
    headers = [
        "Supplier", "Phone", "Address", "Opening Balance",
        "Purchases", "Payments", "Current Balance",
    ]
    ws.append(headers)
    for s in suppliers:
        from models.purchase import Purchase
        from models.payment import Payment
        from sqlalchemy import func
        from extensions import db

        purchases = db.session.query(
            func.coalesce(func.sum(Purchase.subtotal), 0)
        ).filter(Purchase.supplier_id == s.id).scalar()
        payments = db.session.query(
            func.coalesce(func.sum(Payment.amount), 0)
        ).filter(Payment.supplier_id == s.id, Payment.payment_type == "paid").scalar()
        ws.append([
            s.name, s.phone or "", s.address or "",
            s.opening_balance, purchases, payments, s.balance,
        ])
    _style_header(ws)
    return _send_workbook(wb, "suppliers.xlsx")


def export_bills(bills):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bills"
    headers = [
        "Invoice No", "Date", "Customer", "Phone", "Product", "Quantity",
        "Rate", "Discount", "GST", "Total", "Paid", "Balance", "Payment Mode",
    ]
    ws.append(headers)
    for bill in bills:
        for item in bill.items:
            ws.append([
                bill.invoice_no,
                bill.date.strftime("%Y-%m-%d"),
                bill.customer.name if bill.customer else "",
                bill.customer.phone if bill.customer else "",
                item.product.name if item.product else "",
                item.quantity,
                item.price,
                bill.discount,
                bill.gst,
                item.total,
                bill.paid_amount,
                bill.balance,
                bill.payment_mode,
            ])
    _style_header(ws)
    return _send_workbook(wb, "bills.xlsx")


def export_purchases(purchases):
    wb = Workbook()
    ws = wb.active
    ws.title = "Purchases"
    headers = [
        "Purchase ID", "Date", "Supplier", "Product", "Quantity",
        "Purchase Price", "Total", "Paid", "Balance",
    ]
    ws.append(headers)
    for purchase in purchases:
        for item in purchase.items:
            ws.append([
                purchase.id,
                purchase.date.strftime("%Y-%m-%d"),
                purchase.supplier.name if purchase.supplier else "",
                item.product.name if item.product else "",
                item.quantity,
                item.price,
                item.total,
                purchase.paid_amount,
                purchase.balance,
            ])
    _style_header(ws)
    return _send_workbook(wb, "purchases.xlsx")


def export_report(data, headers, title, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for row in data:
        ws.append(row)
    _style_header(ws)
    return _send_workbook(wb, filename)
